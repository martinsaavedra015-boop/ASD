"""Genera el archivo final XLSX (hojas ORDEN + NOTAS) y CSV a partir de las
filas ya armadas por procesar_factura.py / orden_builder.py."""
import os
import csv
import openpyxl
from openpyxl.styles import Font, Alignment
from config import OUTPUT_DIR

ENCABEZADOS = [
    "N", "POSICION ARANCELARIA", "ACUERDO", "ITEM/DESCRIPCION", "COD UNIDAD",
    "MONTO FOB", "CANTIDAD UNIDAD", "CANTIDAD ESTADISTICA", "PESO BRUTO",
    "NUEVO/USADO", "PAIS ORIGEN", "PAIS PROCEDENCIA", "PESO NETO",
    "MARCA LIBRE", "NOMBRE MARCA",
]
COL_ORDEN = list("ABCDEFGHIJKLMNO")


def generar_orden_notas(filas, notas, nombre_archivo):
    """filas: lista de dicts columna->valor (de procesar_factura.py).
    nombre_archivo: solo el nombre, se guarda en OUTPUT_DIR."""
    ruta_salida = os.path.join(OUTPUT_DIR, nombre_archivo)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ORDEN"
    ws.append(ENCABEZADOS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
    for fila in filas:
        ws.append([fila[c] for c in COL_ORDEN])
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.font = Font(name="Arial", size=cell.font.size or 10, bold=cell.font.bold)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=8):
        for cell in row:
            cell.number_format = "@"

    ws2 = wb.create_sheet("NOTAS")
    ws2.append(["#", "Observación"])
    ws2["A1"].font = Font(name="Arial", bold=True)
    ws2["B1"].font = Font(name="Arial", bold=True)
    for i, nota in enumerate(notas, start=1):
        ws2.append([i, nota])
        ws2[f"B{i+1}"].alignment = Alignment(wrap_text=True)

    for ws_ in (ws, ws2):
        for col_cells in ws_.columns:
            length = max((len(str(c.value)) if c.value else 0) for c in col_cells)
            ws_.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 60)

    wb.save(ruta_salida)
    return ruta_salida


def generar_csv(filas, nombre_archivo):
    """Genera el CSV para carga directa en KitApp (DNA), segun el formato
    documentado en Formato_Archivo_Csv.pdf:
    - delimitador ',' (coma), sin fila de encabezado.
    - decimales con punto, nunca coma (la coma es el delimitador de columnas).
    - filas de ITEM: las 15 columnas A-O.
    - filas de SUBITEM: solo 6 columnas (A-F), sin rellenar el resto en blanco.
    - sin comillas en el texto (no debe haber comas dentro de los campos).
    - los campos numericos del item (peso bruto/neto en columnas I y M) nunca
      se dejan vacios -- el validador de la DNA rompe con un campo numerico
      vacio -- se completan con "0.00" cuando la factura no trae el dato
      (queda igual aclarado en NOTAS para completar con packing list/B-L)."""
    ruta_salida = os.path.join(OUTPUT_DIR, nombre_archivo)
    CAMPOS_NUMERICOS_ITEM = {"I", "M"}
    with open(ruta_salida, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=",", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        for fila in filas:
            es_item = fila.get("A") == "N"
            columnas = COL_ORDEN if es_item else COL_ORDEN[:6]
            valores = []
            for c in columnas:
                v = fila.get(c)
                if isinstance(v, str):
                    v = v.replace(",", ".").replace('"', "").replace("'", "")
                if (v is None or v == "") and es_item and c in CAMPOS_NUMERICOS_ITEM:
                    v = "0.00"
                valores.append(v if v is not None else "")
            writer.writerow(valores)
    return ruta_salida
